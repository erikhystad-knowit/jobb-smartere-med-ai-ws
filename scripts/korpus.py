# /// script
# requires-python = ">=3.11"
# dependencies = ["pypdf", "python-docx", "python-pptx", "openpyxl", "extract-msg"]
# ///
"""Inventar og dekningskontroll for store case-korpus.

    uv run scripts/korpus.py inventar <case>   # raw/caser/<case>/ -> korpus/<case>/
    uv run scripts/korpus.py dekning  <case>   # hva er ikke triagert / ikke sitert
    uv run scripts/korpus.py katalog  <case>   # skriver wiki/caser/<case>/katalog.md

Leser kun fra raw/. Skriver til korpus/<case>/ (manifest.csv + tekst/) og,
for katalog, wiki/caser/<case>/katalog.md. triage.csv opprettes tom, men
overskrives aldri.
ID-er (E-001, D-001) og tråd-ID-er (T-001) er stabile mellom kjøringer:
eksisterende manifest leses inn og nye elementer får neste ledige nummer.
"""

import argparse
import csv
import email
import email.policy
import hashlib
import html
import mailbox
import re
import sys
from datetime import datetime
from email.utils import parsedate_to_datetime
from pathlib import Path

ROT = Path(__file__).resolve().parent.parent

EPOST_EXT = {".eml", ".msg", ".mbox"}
TEKST_EXT = {".txt", ".md", ".csv", ".html", ".htm"}
# Når samme dokument finnes i flere formater (OPS-005.csv + OPS-005.xlsx), leses tekst fra det første.
FORMAT_PRIORITET = [".md", ".txt", ".docx", ".pdf", ".pptx", ".csv", ".xlsx", ".html", ".htm"]
# Korpusets egne ID-er i filnavn (EM-001_..., OPS-005_...) og trådmapper (TH-001/) brukes når de finnes.
NATIV_ID = re.compile(r"^([A-Z]+-\d+)(?=[_\-. ]|$)")
TRAD_MAPPE = re.compile(r"^TH-\d+$")
MD_HODE = re.compile(r"^\*\*(Fra|Til|Kopi|Sendt|Dato|Emne):\*\*\s*(.*)$")

MANIFEST_FELT = [
    "id", "type", "mappe", "kilde", "dato", "fra", "til", "kopi", "emne",
    "trad", "trad_nokkel", "vedlegg", "ord", "sha256", "duplikat_av", "feil",
]
TRIAGE_FELT = ["id", "relevans", "tema", "personer", "oppsummering", "status"]

EMNE_PREFIKS = re.compile(r"^\s*((re|sv|vs|fw|fwd|aw|wg|tr)\s*(\[\d+\])?\s*:\s*)+", re.I)


# ---------- tekstuttrekk ----------

def html_til_tekst(s: str) -> str:
    s = re.sub(r"(?is)<(script|style).*?</\1>", "", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</tr>|</li>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = html.unescape(s)
    return re.sub(r"\n\s*\n\s*\n+", "\n\n", s).strip()


def dokument_tekst(navn: str, data: bytes) -> tuple[str, dict]:
    """Returnerer (tekst, metadata) for et dokument. Kaster ved feil."""
    ext = Path(navn).suffix.lower()
    meta: dict = {}
    if ext in TEKST_EXT:
        t = data.decode("utf-8", errors="replace")
        return (html_til_tekst(t) if ext in {".html", ".htm"} else t), meta
    from io import BytesIO
    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(BytesIO(data))
        info = r.metadata or {}
        meta["forfatter"] = str(info.get("/Author") or "")
        meta["tittel"] = str(info.get("/Title") or "")
        meta["dato"] = pdf_dato(str(info.get("/ModDate") or info.get("/CreationDate") or ""))
        sider = [f"[side {i + 1}]\n{(p.extract_text() or '').strip()}" for i, p in enumerate(r.pages)]
        return "\n\n".join(sider), meta
    if ext == ".docx":
        import docx
        d = docx.Document(BytesIO(data))
        cp = d.core_properties
        meta["forfatter"] = cp.author or ""
        meta["tittel"] = cp.title or ""
        meta["dato"] = (cp.modified or cp.created).isoformat() if (cp.modified or cp.created) else ""
        deler = []
        for p in d.paragraphs:
            if p.text.strip():
                stil = p.style.name if p.style is not None else ""
                pre = "#" * int(stil[-1]) + " " if stil.startswith("Heading") and stil[-1:].isdigit() else ""
                deler.append(pre + p.text)
        for t in d.tables:
            for rad in t.rows:
                deler.append(" | ".join(c.text.strip() for c in rad.cells))
        return "\n".join(deler), meta
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(BytesIO(data))
        deler = []
        for i, lysbilde in enumerate(prs.slides, 1):
            tekst = [sh.text_frame.text for sh in lysbilde.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
            deler.append(f"[lysbilde {i}]\n" + "\n".join(tekst))
        return "\n\n".join(deler), meta
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        deler = []
        for ark in wb.worksheets:
            deler.append(f"[ark: {ark.title}]")
            for rad in ark.iter_rows(values_only=True):
                if any(v is not None for v in rad):
                    deler.append(" | ".join("" if v is None else str(v) for v in rad))
        return "\n".join(deler), meta
    raise ValueError(f"ukjent format {ext}")


def pdf_dato(s: str) -> str:
    m = re.match(r"D:(\d{4})(\d{2})(\d{2})", s)
    return f"{m[1]}-{m[2]}-{m[3]}" if m else ""


# ---------- e-post ----------

def les_eposter(sti: Path, rel: str):
    """Gir (kilde, felter, tekst, vedlegg[(navn, data)]) for hver melding i fila."""
    ext = sti.suffix.lower()
    if ext == ".msg":
        yield rel, *les_msg(sti)
    elif ext == ".mbox":
        for i, m in enumerate(mailbox.mbox(str(sti), factory=None)):
            msg = email.message_from_bytes(m.as_bytes(), policy=email.policy.default)
            yield f"{rel}#{i + 1}", *les_eml(msg)
    else:
        msg = email.message_from_bytes(sti.read_bytes(), policy=email.policy.default)
        yield rel, *les_eml(msg)


def er_md_epost(sti: Path) -> bool:
    if sti.suffix.lower() not in {".md", ".txt"}:
        return False
    hode = sti.read_text(encoding="utf-8", errors="replace").splitlines()[:15]
    return any(MD_HODE.match(linje) and linje.startswith("**Fra:**") for linje in hode)


def les_md_epost(sti: Path):
    """E-post eksportert som markdown: '# Emne', så '**Fra:** …'-linjer, så brødtekst."""
    navn = {"Fra": "fra", "Til": "til", "Kopi": "kopi", "Sendt": "dato", "Dato": "dato", "Emne": "emne"}
    felt = {k: "" for k in ("dato", "fra", "til", "kopi", "emne", "message_id", "in_reply_to", "references")}
    linjer = sti.read_text(encoding="utf-8", errors="replace").splitlines()
    i = 0
    while i < len(linjer) and (not linjer[i].strip() or linjer[i].startswith("# ") or MD_HODE.match(linjer[i])):
        m = MD_HODE.match(linjer[i])
        if m:
            felt[navn[m[1]]] = m[2].strip()
        i += 1
    return felt, "\n".join(linjer[i:]), []


def les_eml(msg):
    def h(navn):
        v = msg.get(navn)
        return str(v).strip() if v is not None else ""

    dato = ""
    try:
        if h("Date"):
            dato = parsedate_to_datetime(h("Date")).isoformat()
    except (TypeError, ValueError):
        pass
    felt = {
        "dato": dato, "fra": h("From"), "til": h("To"), "kopi": h("Cc"), "emne": h("Subject"),
        "message_id": h("Message-ID"), "in_reply_to": h("In-Reply-To"), "references": h("References"),
    }
    kropp = msg.get_body(preferencelist=("plain", "html"))
    tekst = ""
    if kropp is not None:
        tekst = kropp.get_content()
        if kropp.get_content_type() == "text/html":
            tekst = html_til_tekst(tekst)
    vedlegg = []
    for del_ in msg.iter_attachments():
        navn = del_.get_filename() or "uten-navn"
        data = del_.get_payload(decode=True) or b""
        if del_.get_content_type() == "message/rfc822":
            navn = navn if navn != "uten-navn" else "videresendt.eml"
            data = del_.get_content().as_bytes() if hasattr(del_.get_content(), "as_bytes") else data
        vedlegg.append((navn, data))
    return felt, tekst, vedlegg


def les_msg(sti: Path):
    import extract_msg
    m = extract_msg.Message(str(sti))
    try:
        dato = m.date.isoformat() if isinstance(m.date, datetime) else str(m.date or "")
        felt = {
            "dato": dato, "fra": m.sender or "", "til": m.to or "", "kopi": m.cc or "",
            "emne": m.subject or "",
            "message_id": getattr(m, "messageId", "") or "",
            "in_reply_to": getattr(m, "inReplyTo", "") or "",
            "references": "",
        }
        tekst = m.body or ""
        if not tekst.strip() and getattr(m, "htmlBody", None):
            tekst = html_til_tekst(m.htmlBody.decode("utf-8", errors="replace"))
        vedlegg = []
        for a in m.attachments:
            data = a.data if isinstance(a.data, bytes) else b""
            vedlegg.append((a.longFilename or a.shortFilename or "uten-navn", data))
        return felt, tekst, vedlegg
    finally:
        m.close()


def trad_nokkel(felt: dict, rot_for: dict, emne_for: dict) -> str:
    """Trådnøkkel: eldste kjente Message-ID i kjeden, ellers tråden med samme normaliserte emne."""
    emne = EMNE_PREFIKS.sub("", felt["emne"]).strip().lower()
    refs = re.findall(r"<[^>]+>", felt["references"]) or re.findall(r"<[^>]+>", felt["in_reply_to"])
    if refs:
        nk = rot_for.get(refs[0], refs[0])
    else:
        nk = emne_for.get(emne) or felt["message_id"] or "emne:" + emne
    if emne:
        emne_for.setdefault(emne, nk)
    return nk


# ---------- inventar ----------

def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def les_csv(sti: Path) -> list[dict]:
    if not sti.exists():
        return []
    with sti.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def neste_id(brukt: set[str], prefiks: str) -> str:
    n = 1 + max((int(i[len(prefiks):]) for i in brukt if re.fullmatch(re.escape(prefiks) + r"\d+", i)), default=0)
    return f"{prefiks}{n:03d}"


def inventar(case: str, rot: Path, raw: Path | None = None) -> int:
    rawdir = (rot / raw if raw else rot / "raw" / "caser" / case).resolve()
    ut = rot / "korpus" / case
    if not rawdir.is_dir():
        print(f"Fant ikke {rawdir}", file=sys.stderr)
        return 1
    (ut / "tekst").mkdir(parents=True, exist_ok=True)

    gammel = {r["kilde"]: r for r in les_csv(ut / "manifest.csv")}
    brukte_id = {r["id"] for r in gammel.values()}
    gammel_trad = {r["trad_nokkel"]: r["trad"] for r in gammel.values() if r.get("trad")}

    filer = sorted(
        p for p in rawdir.rglob("*")
        if p.is_file() and not p.name.startswith(".") and "__MACOSX" not in p.parts
    )
    poster, tekster, vedlegg_hash = [], {}, {}
    dokgrupper: dict[str, list[Path]] = {}

    def mappe(sti: Path) -> str:
        deler = sti.relative_to(rawdir).parts
        return deler[0] if len(deler) > 1 else ""

    for sti in filer:
        rel = str(sti.relative_to(rot))
        nativ = NATIV_ID.match(sti.stem)
        if sti.suffix.lower() not in EPOST_EXT and not er_md_epost(sti):
            dokgrupper.setdefault(nativ[1] if nativ else rel, []).append(sti)
            continue
        try:
            meldinger = [(rel, *les_md_epost(sti))] if sti.suffix.lower() in {".md", ".txt"} else list(les_eposter(sti, rel))
        except Exception as e:  # noqa: BLE001 — rapporteres i manifestet
            meldinger = [(rel, {k: "" for k in ("dato", "fra", "til", "kopi", "emne", "message_id", "in_reply_to", "references")}, "", [])]
            feil = f"{type(e).__name__}: {e}"
        else:
            feil = ""
        for kilde, felt, tekst, vedlegg in meldinger:
            pid = gammel.get(kilde, {}).get("id") or (nativ[1] if nativ and len(meldinger) == 1 else neste_id(brukte_id, "E-"))
            brukte_id.add(pid)
            deler = [tekst.strip()]
            for navn, data in vedlegg:
                vedlegg_hash.setdefault(sha(data), []).append(pid)
                try:
                    vt, _ = dokument_tekst(navn, data)
                    deler.append(f"\n\n---\n## Vedlegg: {navn}\n\n{vt.strip()}")
                except Exception as e:  # noqa: BLE001
                    deler.append(f"\n\n---\n## Vedlegg: {navn}\n\n_(ikke uttrukket: {e})_")
            post = {
                "id": pid, "type": "epost", "mappe": mappe(sti), "kilde": kilde, "dato": felt["dato"],
                "fra": felt["fra"], "til": felt["til"], "kopi": felt["kopi"], "emne": felt["emne"],
                "vedlegg": "; ".join(n for n, _ in vedlegg), "sha256": sha((felt["message_id"] + tekst).encode()),
                "feil": feil, "_felt": felt,
            }
            if TRAD_MAPPE.match(sti.parent.name):
                post["trad"] = sti.parent.name
                post["trad_nokkel"] = "mappe:" + sti.parent.name
            poster.append(post)
            tekster[pid] = "".join(deler)

    for nokkel, stier in dokgrupper.items():
        stier.sort(key=lambda p: FORMAT_PRIORITET.index(p.suffix.lower()) if p.suffix.lower() in FORMAT_PRIORITET else 99)
        kilde = "; ".join(str(p.relative_to(rot)) for p in stier)
        pid = gammel.get(kilde, {}).get("id") or (nokkel if NATIV_ID.match(nokkel) else neste_id(brukte_id, "D-"))
        brukte_id.add(pid)
        tekst, meta, feil = "", {}, ""
        for sti in stier:  # første format som gir tekst vinner
            try:
                tekst, meta = dokument_tekst(sti.name, sti.read_bytes())
            except Exception as e:  # noqa: BLE001
                feil = f"{sti.suffix}: {type(e).__name__}: {e}"
                continue
            if tekst.strip():
                feil = ""
                break
            feil = "ingen tekst (skannet? trenger OCR)"
        if len({p.stem for p in stier}) > 1:
            feil = "; ".join(filter(None, [feil, "ulike filnavn under samme ID — sjekk om det er versjoner"]))
        poster.append({
            "id": pid, "type": "dokument", "mappe": mappe(stier[0]), "kilde": kilde, "dato": meta.get("dato", ""),
            "fra": meta.get("forfatter", ""), "til": "", "kopi": "",
            "emne": meta.get("tittel") or stier[0].stem, "vedlegg": "", "sha256": sha(stier[0].read_bytes()), "feil": feil,
        })
        tekster[pid] = tekst

    # Tråder: rot-Message-ID for hver melding, så stabile T-ID-er sortert på første dato.
    rot_for, emne_for = {}, {}
    for p in sorted((p for p in poster if p["type"] == "epost" and not p.get("trad")), key=lambda p: p["dato"]):
        nk = trad_nokkel(p["_felt"], rot_for, emne_for)
        if p["_felt"]["message_id"]:
            rot_for[p["_felt"]["message_id"]] = nk
        p["trad_nokkel"] = nk
    brukte_trad = set(gammel_trad.values())
    for p in sorted((p for p in poster if p["type"] == "epost" and not p.get("trad")), key=lambda p: p["dato"]):
        nk = p["trad_nokkel"]
        if nk not in gammel_trad:
            gammel_trad[nk] = neste_id(brukte_trad, "T-")
            brukte_trad.add(gammel_trad[nk])
        p["trad"] = gammel_trad[nk]

    # Duplikater: dokumenter som også ligger som e-postvedlegg, og identiske filer.
    sett = {}
    for p in poster:
        if p["type"] == "dokument":
            treff = vedlegg_hash.get(p["sha256"], []) + ([sett[p["sha256"]]] if p["sha256"] in sett else [])
            p["duplikat_av"] = "; ".join(treff)
            sett.setdefault(p["sha256"], p["id"])

    poster.sort(key=lambda p: (p["type"], p["id"]))
    for p in poster:
        hode = "\n".join(f"{k}: {p.get(k, '')}" for k in ("id", "mappe", "kilde", "dato", "fra", "til", "kopi", "emne", "trad", "vedlegg") if p.get(k))
        p["ord"] = len(tekster[p["id"]].split())
        (ut / "tekst" / f"{p['id']}.md").write_text(f"```\n{hode}\n```\n\n{tekster[p['id']]}\n", encoding="utf-8")

    with (ut / "manifest.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=MANIFEST_FELT, extrasaction="ignore")
        w.writeheader()
        w.writerows(poster)

    triage = ut / "triage.csv"
    if not triage.exists():
        with triage.open("w", newline="", encoding="utf-8") as f:
            csv.DictWriter(f, fieldnames=TRIAGE_FELT).writeheader()

    eposter = [p for p in poster if p["type"] == "epost"]
    dok = [p for p in poster if p["type"] == "dokument"]
    datoer = sorted(p["dato"][:10] for p in poster if p["dato"])
    print(f"Inventar for '{case}' -> {ut.relative_to(rot)}")
    print(f"  e-poster:   {len(eposter)} i {len({p['trad'] for p in eposter})} tråder")
    print(f"  dokumenter: {len(dok)}  (duplikat av vedlegg/annen fil: {sum(1 for p in dok if p['duplikat_av'])})")
    if datoer:
        print(f"  periode:    {datoer[0]} – {datoer[-1]}")
    print(f"  ord totalt: {sum(p['ord'] for p in poster):,}".replace(",", " "))
    nye = [p["id"] for p in poster if p["kilde"] not in gammel]
    print(f"  nye siden forrige kjøring: {len(nye)}")
    feilede = [p for p in poster if p["feil"]]
    for p in feilede:
        print(f"  ! {p['id']} {p['kilde']}: {p['feil']}")
    return 0


# ---------- dekning ----------

def dekning(case: str, rot: Path) -> int:
    ut = rot / "korpus" / case
    manifest = les_csv(ut / "manifest.csv")
    if not manifest:
        print("Ingen manifest — kjør inventar først.", file=sys.stderr)
        return 1
    triage = {r["id"]: r for r in les_csv(ut / "triage.csv")}
    wikidir = rot / "wiki" / "caser" / case
    sitert = set()
    for md in wikidir.rglob("*.md") if wikidir.is_dir() else []:
        if md.name == "katalog.md":  # generert, lister alt — teller ikke som sitering
            continue
        sitert |= set(re.findall(r"\b[A-Z]+-\d+\b", md.read_text(encoding="utf-8")))

    ids = [r["id"] for r in manifest]
    utriagert = [i for i in ids if i not in triage]
    viktige = [i for i in ids if triage.get(i, {}).get("relevans") in ("hoy", "middels")]
    hoy_usitert = [i for i in ids if triage.get(i, {}).get("relevans") == "hoy" and i not in sitert]
    middels_usitert = [i for i in ids if triage.get(i, {}).get("relevans") == "middels" and i not in sitert]

    print(f"Dekning for '{case}': {len(ids)} elementer")
    print(f"  triagert:            {len(ids) - len(utriagert)}/{len(ids)}")
    print(f"  høy/middels:         {len(viktige)}")
    print(f"  sitert i wikien:     {len(sitert & set(ids))}")
    print(f"  HØY, ikke sitert:    {len(hoy_usitert)}  {' '.join(hoy_usitert)}")
    print(f"  middels, ikke sitert:{len(middels_usitert):>3}  {' '.join(middels_usitert[:40])}{' …' if len(middels_usitert) > 40 else ''}")
    if utriagert:
        print(f"  utriagert: {' '.join(utriagert[:60])}{' …' if len(utriagert) > 60 else ''}")
    ukjente = sorted(set(triage) - set(ids))
    if ukjente:
        print(f"  ! triage-rader uten manifest-treff: {' '.join(ukjente)}")
    return 0 if not (utriagert or hoy_usitert) else 2


# ---------- katalog ----------

def katalog(case: str, rot: Path) -> int:
    """Skriver wiki/caser/<case>/katalog.md: alle elementer gruppert på tema og tråd."""
    ut = rot / "korpus" / case
    manifest = les_csv(ut / "manifest.csv")
    if not manifest:
        print("Ingen manifest — kjør inventar først.", file=sys.stderr)
        return 1
    triage = {r["id"]: r for r in les_csv(ut / "triage.csv")}

    grupper: dict[str, list[dict]] = {}
    stoy = []
    for r in manifest:
        t = triage.get(r["id"], {})
        if t.get("relevans") == "stoy":
            stoy.append(r["id"])
            continue
        tema = (t.get("tema") or "").split(";")[0].strip() or ("_utriagert" if not t else "_uten-tema")
        grupper.setdefault(tema, []).append({**r, **t})

    def celle(s: str) -> str:
        return (s or "").replace("|", "\\|").replace("\n", " ").strip()

    def dato(s: str) -> str:
        return s[:10]

    i_dag = datetime.now().date().isoformat()
    linjer = [
        "---", f"title: Katalog — {case}", "type: korpus", f"case: {case}",
        f"opprettet: {i_dag}", f"oppdatert: {i_dag}", "tags: [katalog, generert]", "---", "",
        f"# Katalog — {case}", "",
        f"> Generert av `scripts/korpus.py katalog {case}`. Ikke rediger for hånd.",
        f"> {len(manifest)} elementer · {len(triage)} triagert · {len(stoy)} støy.",
        f"> Tilbake til [[caser/{case}/oversikt|oversikten]].", "",
    ]
    for tema in sorted(grupper, key=lambda t: (t.startswith("_"), t)):
        rader = grupper[tema]
        overskrift = {"_utriagert": "Utriagert", "_uten-tema": "Uten tema"}.get(tema, f"[[caser/{case}/tema/{tema}|{tema}]]")
        linjer += [f"## {overskrift}", "", "| ID | Dato | Fra | Emne | Tråd | Rel. | Oppsummering |", "|---|---|---|---|---|---|---|"]
        rader.sort(key=lambda r: (r.get("trad") or "~" + r["id"], r["dato"]))
        for r in rader:
            fra = re.sub(r"\s*<[^>]*>", "", r["fra"])
            linjer.append(
                f"| {r['id']} | {dato(r['dato'])} | {celle(fra)[:30]} | {celle(r['emne'])[:60]} | "
                f"{r.get('trad', '')} | {r.get('relevans', '')} | {celle(r.get('oppsummering', ''))} |"
            )
        linjer.append("")
    if stoy:
        linjer += ["## Støy", "", " ".join(stoy), ""]

    mappe = rot / "wiki" / "caser" / case
    mappe.mkdir(parents=True, exist_ok=True)
    (mappe / "katalog.md").write_text("\n".join(linjer), encoding="utf-8")
    print(f"Skrev {(mappe / 'katalog.md').relative_to(rot)} ({len(manifest)} elementer, {len(grupper)} grupper)")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("kommando", choices=["inventar", "dekning", "katalog"])
    ap.add_argument("case", help="case-slug, f.eks. anbud-2026")
    ap.add_argument("--raw", type=Path, help="kilde-mappe relativt til repo-roten (standard: raw/caser/<case>)")
    ap.add_argument("--rot", type=Path, default=ROT, help="repo-rot (standard: mappa over scripts/)")
    a = ap.parse_args()
    rot = a.rot.resolve()
    if a.kommando == "inventar":
        return inventar(a.case, rot, a.raw)
    return {"dekning": dekning, "katalog": katalog}[a.kommando](a.case, rot)


if __name__ == "__main__":
    sys.exit(main())
